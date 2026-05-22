"""
Compare cxt and hyh evaluation CSVs, find disagreements, extract evidence from eval_reports,
generate:
1. disagreement_summary.csv — Paper_ID, Field, cxt_value, hyh_value, Evidence
2. cxt_with_disagreements.xlsx — cxt's table with orange-highlighted disagreement cells

Usage:
    uv run python compare/compare_evaluations.py
"""

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

COMPARE_DIR = Path(__file__).parent
PROJECT_DIR = COMPARE_DIR.parent
CXT_CSV = COMPARE_DIR / "cxt's evaluation.csv"
HYH_CSV = COMPARE_DIR / "hyh's evaluation.csv"
EVAL_REPORTS_DIR = PROJECT_DIR / "eval_reports"
OUTPUT_CSV = COMPARE_DIR / "disagreement_summary.csv"
OUTPUT_XLSX = COMPARE_DIR / "cxt_with_disagreements.xlsx"

# Fields to exclude (metadata or free-text)
EXCLUDE_FIELDS = {
    "Paper_ID", "Citation_Key", "Bibtex", "Title", "Year", "Venue",
    "Domain", "Focus_Type", "Temporal_Modeling_Details", "Raw Eval Metrics",
    "Clinical_Theory", "Agreement Method",
}

# Value normalization
VALUE_NORMALIZE = {
    "FULL DISCLOSURE": "FULL",
    "PARTIAL DISCLOSURE": "PARTIAL",
    "NO DISCLOSURE": "NO",
}

# Column name fix for cxt typo
COL_RENAME_CXT = {"Interention_Sensitivity": "Intervention_Sensitivity"}

# Mapping: CSV column name -> eval_report ### header field name
CSV_TO_REPORT = {
    "Simulation_Target": "simulation_target",
    "Persona_Model_Depth": "persona_model_depth",
    "Uses_Dynamic_State": "uses_dynamic_state",
    "Eval_Human_Experts": "eval_human_experts",
    "Eval_Lay_Users": "eval_lay_users",
    "Eval_User_Study": "eval_user_study",
    "Eval_LLM_Judge": "eval_llm_judge",
    "Eval_Automatic": "eval_automatic",
    "Safety": "dim_safety",
    "Realism": "dim_realism",
    "Consistency": "dim_consistency",
    "Fidelity": "dim_fidelity",
    "Human Learning / Outcomes": "dim_human_learning",
    "Utility": "dim_utility",
    "Emotional Plausibility": "dim_emotional_plausibility",
    "Theory_Operationalized": "theory_operationalized",
    "Behavior_Eval_Depth": "behavior_eval_depth",
    "Intervention_Sensitivity": "intervention_sensitivity",
    "Interaction_Level": "interaction_level",
    "Prompt_Disclosure": "prompt_disclosure",
    "Theory_Grounding": "theory_grounding",
    "Reliability_Reported": "reliability_reported",
    "Coding Options": "coding_options",
    "Has_Rubric": "has_rubric",
    "LLM_Judge_Validated": "llm_judge_validated",
    "Uses_Standard_Metrics": "uses_standard_metrics",
    "Metric_Interpretable": "metric_interpretable",
    "Comparable_To_Prior_Work": "comparable_to_prior_work",
    "Has_Longitudinal_Eval": "has_longitudinal_eval",
    "Has_Robustness_Testing": "has_robustness_testing",
    "Has_Failure_Analysis": "has_failure_analysis",
    "Sim_Behavior_Realistic": "sim_behavior_realistic",
    "Dataset_Available": "dataset_available",
}


def norm_pid(pid: str) -> str:
    if pd.isna(pid):
        return ""
    return re.sub(r"[（(]综述[）)]", "", str(pid).strip()).strip()


def load_eval_report(pid: str) -> str | None:
    """Load eval report markdown for a given paper ID."""
    report_path = EVAL_REPORTS_DIR / f"{pid}.md"
    if report_path.exists():
        return report_path.read_text(encoding="utf-8")
    return None


def extract_evidence(report_text: str, field_name: str) -> str:
    """Extract evidence from eval report for a given field.

    Looks for ### field_name: ... section and extracts the Evidence line.
    """
    report_field = CSV_TO_REPORT.get(field_name)
    if not report_field:
        return ""

    # Find the section header
    pattern = rf"###\s+{re.escape(report_field)}\s*:\s*.*?(?:\(confidence:\s*\d+\))?\s*\n(.*?)(?=\n###|\n##|\Z)"
    match = re.search(pattern, report_text, re.DOTALL | re.IGNORECASE)
    if not match:
        return ""

    section = match.group(1).strip()

    # Extract evidence line(s)
    evidence_lines = []
    for line in section.split("\n"):
        line = line.strip()
        if line.startswith("**Evidence**"):
            # Remove the **Evidence** (p.X): prefix
            cleaned = re.sub(r"^\*\*Evidence\*\*\s*(?:\([^)]*\))?\s*:?\s*", "", line)
            evidence_lines.append(cleaned)
        elif line.startswith("- **") and "Evidence" in line:
            cleaned = re.sub(r"^-\s*\*\*Evidence\*\*\s*(?:\([^)]*\))?\s*:?\s*", "", line)
            evidence_lines.append(cleaned)

    if evidence_lines:
        return " | ".join(evidence_lines)

    # Fallback: return the whole section (truncated)
    return section[:300] + "..." if len(section) > 300 else section


def main():
    # Load CSVs
    cxt = pd.read_csv(CXT_CSV, dtype=str).dropna(subset=["Paper_ID"])
    hyh = pd.read_csv(HYH_CSV, dtype=str).dropna(subset=["Paper_ID"])

    # Normalize PIDs
    cxt["_pid"] = cxt["Paper_ID"].apply(norm_pid)
    hyh["_pid"] = hyh["Paper_ID"].apply(norm_pid)

    # Exclude 综述
    cxt = cxt[~cxt["Paper_ID"].str.contains("综述", na=False)].copy()
    hyh = hyh[~hyh["Paper_ID"].str.contains("综述", na=False)].copy()

    # Merge
    merged = cxt.merge(hyh, on="_pid", suffixes=("_cxt", "_hyh"))
    print(f"Merged: {len(merged)} papers")

    # Find comparable fields
    fields = set()
    for col in merged.columns:
        if col.endswith("_cxt"):
            base = col[:-4]
            if base not in EXCLUDE_FIELDS and f"{base}_hyh" in merged.columns:
                fields.add(base)

    # Find disagreements and extract evidence
    disagreements = []
    for field in sorted(fields):
        col_cxt = f"{field}_cxt"
        col_hyh = f"{field}_hyh"
        df = merged[["_pid", col_cxt, col_hyh]].copy()
        df["v_cxt"] = df[col_cxt].str.strip().str.upper().replace(VALUE_NORMALIZE)
        df["v_hyh"] = df[col_hyh].str.strip().str.upper().replace(VALUE_NORMALIZE)

        mask = (
            (df["v_cxt"] != df["v_hyh"])
            & df["v_cxt"].notna()
            & df["v_hyh"].notna()
            & (df["v_cxt"] != "")
            & (df["v_hyh"] != "")
        )
        disagree_df = df[mask]

        for _, row in disagree_df.iterrows():
            pid = row["_pid"]
            report = load_eval_report(pid)
            evidence = ""
            if report:
                evidence = extract_evidence(report, field)

            disagreements.append({
                "Paper_ID": pid,
                "Field": field,
                "cxt_value": row["v_cxt"],
                "hyh_value": row["v_hyh"],
                "Evidence": evidence,
            })

    # Sort by Paper_ID (numeric)
    disagree_df = pd.DataFrame(disagreements)
    disagree_df = disagree_df.sort_values(
        ["Paper_ID", "Field"],
        key=lambda x: pd.to_numeric(x, errors="coerce").fillna(999),
    )

    # Save summary CSV
    disagree_df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"Saved: {OUTPUT_CSV} ({len(disagree_df)} disagreements)")

    # Generate xlsx based on cxt's table
    # First save cxt as xlsx
    cxt_clean = cxt.drop(columns=["_pid"])
    cxt_clean.to_excel(OUTPUT_XLSX, index=False, sheet_name="cxt evaluation")

    # Open and format with openpyxl
    from openpyxl import load_workbook
    wb = load_workbook(OUTPUT_XLSX)
    ws = wb["cxt evaluation"]

    # Build lookup of disagreement cells: (paper_id, col_name) -> True
    disagree_lookup = set()
    for _, row in disagree_df.iterrows():
        disagree_lookup.add((str(row["Paper_ID"]), row["Field"]))

    # Find column indices
    header_row = [cell.value for cell in ws[1]]
    col_index = {}
    for idx, name in enumerate(header_row):
        if name:
            col_index[name.strip()] = idx + 1  # 1-based

    # Find Paper_ID column
    pid_col = col_index.get("Paper_ID", 1)

    # Orange fill
    orange_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")

    # Apply orange to disagreement cells
    count = 0
    for row_idx in range(2, ws.max_row + 1):
        pid = str(ws.cell(row_idx, pid_col).value or "").strip()
        pid = norm_pid(pid)
        for field_name in disagree_lookup:
            paper_id, field = field_name
            if pid == paper_id and field in col_index:
                cell = ws.cell(row_idx, col_index[field])
                cell.fill = orange_fill
                count += 1

    wb.save(OUTPUT_XLSX)
    print(f"Saved: {OUTPUT_XLSX} ({count} cells highlighted orange)")

    # Print summary
    print(f"\nTotal disagreements: {len(disagree_df)}")
    print(f"Papers with disagreements: {disagree_df['Paper_ID'].nunique()}")
    print(f"Fields with disagreements: {disagree_df['Field'].nunique()}")


if __name__ == "__main__":
    main()
