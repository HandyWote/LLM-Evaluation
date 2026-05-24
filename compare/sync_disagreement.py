"""
Sync corrected values from disagreement_cn.csv to cxt_with_disagreements.xlsx.
- Modified cells get red fill
- Sheet2 gets modification reasons
"""

import csv
import openpyxl
from openpyxl.styles import PatternFill

# Value normalization maps: CSV -> XLSX
BOOL_MAP = {"YES": "Yes", "NO": "No"}

BEHAVIOR_DEPTH_MAP = {
    "DYNAMIC": "Dynamic",
    "STATIC": "Static",
    "PATTERN-LEVEL": "Pattern-level",
    "NONE": "None",
}

INTERACTION_LEVEL_MAP = {
    "EXTENDED DIALOGUE": "Extended Dialogue",
    "SINGLE-TURN": "Single-turn",
    "SHORT MULTI-TURN": "Short Multi-turn",
    "LONGITUDINAL": "Longitudinal",
}

PROMPT_DISCLOSURE_MAP = {
    "FULL": "Full Disclosure",
    "PARTIAL": "Partial Disclosure",
    "NO": "No Disclosure",
}

PERSONA_MODEL_MAP = {
    "SURFACE PERSONA": "Surface Persona",
    "BEHAVIORAL STATE MODEL": "Behavioral State Model",
    "COGNITIVE MODEL": "Cognitive Model",
    "DYNAMIC TRAITS": "Dynamic Traits",
    "EXPERT PRINCIPLES": "Expert Principles",
    # compound values
    "EXPERT PRINCIPLES/BEHAVIORAL MODEL": "Expert Principles",
    "SURFACE PERSONA/DYNAMIC TRAITS": "Surface Persona/Dynamic Traits",
    "COGNITIVE MODEL/EXPERT PRINCIPLES": "Cognitive Model/Expert Principles",
    "COGNITIVE MODEL / BEHAVIORAL STATE": "Cognitive Model",
    "EXPERT PRINCIPLES/SURFACE PERSONA": "Expert Principles",
    "COGNITIVE MODEL/DYNAMIC TRAITS": "Cognitive Model/Dynamic Traits",
}

SIMULATION_TARGET_MAP = {
    "CLIENT AGENT": "Client Agent",
    "THERAPIST AGENT": "Therapist Agent",
    "DUAL-AGENT": "Dual-agent",
    "HUMAN TRAINEE": "Human Trainee",
    "PERSONA AGENT": "Persona Agent",
    "HUMAN TRAINEE/DUAL-AGENT": "Human Trainee/Dual-agent",
    "HUMAN TRAINEE/REFRAMING ASSISTANT": "Human Trainee",
}

THEORY_OP_MAP = {
    "STRONG": "Strong",
    "PARTIAL": "Partial",
    "MENTIONED": "Mentioned",
    "NONE": "None",
}

THEORY_GRND_MAP = {
    "STRONG": "Strong",
    "WEAK": "Weak",
    "NONE": "None",
}

# Map field names to their value normalizers
FIELD_NORMALIZERS = {
    "Eval_Automatic": BOOL_MAP,
    "Eval_Lay_Users": BOOL_MAP,
    "Eval_Human_Experts": BOOL_MAP,
    "Eval_LLM_Judge": BOOL_MAP,
    "Eval_User_Study": BOOL_MAP,
    "Safety": BOOL_MAP,
    "Realism": BOOL_MAP,
    "Consistency": BOOL_MAP,
    "Fidelity": BOOL_MAP,
    "Human Learning / Outcomes": BOOL_MAP,
    "Utility": BOOL_MAP,
    "Emotional Plausibility": BOOL_MAP,
    "Coding Options": BOOL_MAP,  # YES/NO in CSV
    "Reliability_Reported": BOOL_MAP,
    "LLM_Judge_Validated": BOOL_MAP,
    "Uses_Standard_Metrics": BOOL_MAP,
    "Comparable_To_Prior_Work": BOOL_MAP,
    "Has_Longitudinal_Eval": BOOL_MAP,
    "Has_Robustness_Testing": BOOL_MAP,
    "Has_Failure_Analysis": BOOL_MAP,
    "Has_Rubric": BOOL_MAP,
    "Sim_Behavior_Realistic": BOOL_MAP,
    "Dataset_Available": BOOL_MAP,
    "Uses_Dynamic_State": BOOL_MAP,
    "Behavior_Eval_Depth": BEHAVIOR_DEPTH_MAP,
    "Interaction_Level": INTERACTION_LEVEL_MAP,
    "Prompt_Disclosure": PROMPT_DISCLOSURE_MAP,
    "Persona_Model_Depth": PERSONA_MODEL_MAP,
    "Simulation_Target": SIMULATION_TARGET_MAP,
    "Theory_Operationalized": THEORY_OP_MAP,
    "Theory_Grounding": THEORY_GRND_MAP,
}


def normalize_value(field: str, raw: str) -> str:
    """Convert CSV value to xlsx format."""
    raw = raw.strip()
    normalizer = FIELD_NORMALIZERS.get(field)
    if normalizer:
        # Try exact match first
        if raw in normalizer:
            return normalizer[raw]
        # Try upper-case match
        if raw.upper() in normalizer:
            return normalizer[raw.upper()]
    # Fallback: Title case for simple YES/NO-like
    if raw.upper() in ("YES", "NO"):
        return raw.capitalize()
    return raw


def main():
    csv_path = "compare/disagreement_cn.csv"
    xlsx_path = "compare/cxt_with_disagreements.xlsx"

    # 1. Read CSV corrections
    corrections = []  # (paper_id, field, new_value, reason)
    with open(csv_path, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            pid = int(row["Paper_ID"])
            field = row["字段"].strip()
            new_val = row["修正判定"].strip()
            reason = row["理由"].strip()
            if new_val:  # only rows with a correction
                corrections.append((pid, field, new_val, reason))

    print(f"Found {len(corrections)} corrections to apply")

    # 2. Load xlsx
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["cxt evaluation"]

    # Build header -> column index map
    header_to_col = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            header_to_col[val.strip()] = col

    # Build paper_id -> row index map
    pid_to_row = {}
    for row in range(2, ws.max_row + 1):
        pid = ws.cell(row=row, column=1).value
        if pid is not None:
            pid_to_row[int(pid)] = row

    # 3. Create Sheet2 for reasons
    if "修改理由" in wb.sheetnames:
        del wb["修改理由"]
    ws2 = wb.create_sheet("修改理由")
    ws2.cell(row=1, column=1, value="Paper_ID")
    ws2.cell(row=1, column=2, value="字段")
    ws2.cell(row=1, column=3, value="原值(cxt)")
    ws2.cell(row=1, column=4, value="修正值")
    ws2.cell(row=1, column=5, value="理由")
    # Bold header
    from openpyxl.styles import Font
    bold = Font(bold=True)
    for col in range(1, 6):
        ws2.cell(row=1, column=col).font = bold

    # 4. Apply corrections
    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    applied = 0
    skipped = 0

    for i, (pid, field, new_val_raw, reason) in enumerate(corrections, start=2):
        if pid not in pid_to_row:
            print(f"  SKIP: Paper_ID={pid} not found in xlsx")
            skipped += 1
            continue
        if field not in header_to_col:
            print(f"  SKIP: Field '{field}' not found in xlsx headers")
            skipped += 1
            continue

        row_idx = pid_to_row[pid]
        col_idx = header_to_col[field]
        cell = ws.cell(row=row_idx, column=col_idx)
        old_val = cell.value

        # Normalize new value
        new_val = normalize_value(field, new_val_raw)

        # Apply
        cell.value = new_val
        cell.fill = red_fill

        # Write reason to Sheet2
        ws2.cell(row=i, column=1, value=pid)
        ws2.cell(row=i, column=2, value=field)
        ws2.cell(row=i, column=3, value=old_val)
        ws2.cell(row=i, column=4, value=new_val)
        ws2.cell(row=i, column=5, value=reason)

        applied += 1
        if old_val != new_val:
            print(f"  Paper {pid} | {field}: '{old_val}' -> '{new_val}'")

    # 5. Auto-fit Sheet2 column widths
    for col in range(1, 6):
        max_len = 0
        for row in range(1, ws2.max_row + 1):
            val = ws2.cell(row=row, column=col).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max_len + 2, 80)

    # 6. Save
    wb.save(xlsx_path)
    print(f"\nDone: {applied} applied, {skipped} skipped. Saved to {xlsx_path}")


if __name__ == "__main__":
    main()
